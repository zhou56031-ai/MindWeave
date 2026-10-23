"""PDF、DOCX、PPTX 的文字与页号提取回归。"""

import unittest
from io import BytesIO

import pymupdf
from docx import Document
from fastapi import HTTPException
from pptx import Presentation

from app.services.document_extract import extract_pages
from app.services.pdf_vision import image_path, remove_images


class DocumentExtractTests(unittest.TestCase):
    def test_pdf_page_number(self):
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "PDF page text")
        raw = pdf.tobytes()
        pdf.close()

        safe_name, pages = extract_pages("sample.pdf", raw)

        self.assertEqual(safe_name, "sample.pdf")
        self.assertEqual(pages[0][0], 1)
        self.assertIn("PDF page text", pages[0][1])

    def test_docx_has_no_print_page_number(self):
        doc = Document()
        doc.add_paragraph("Word paragraph")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        buffer = BytesIO()
        doc.save(buffer)

        _, pages = extract_pages("sample.docx", buffer.getvalue())

        self.assertIsNone(pages[0][0])
        self.assertIn("Word paragraph", pages[0][1])
        self.assertIn("A | B", pages[0][1])

    def test_pptx_slide_number(self):
        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[6])
        shape = slide.shapes.add_textbox(0, 0, 3_000_000, 1_000_000)
        shape.text = "PPT slide text"
        buffer = BytesIO()
        slides.save(buffer)

        _, pages = extract_pages("sample.pptx", buffer.getvalue())

        self.assertEqual(pages[0][0], 1)
        self.assertIn("PPT slide text", pages[0][1])

    def test_image_path_rejects_non_page_name(self):
        user_id = "extract-test-user"
        md5_value = "a" * 32
        path = image_path(user_id, md5_value, "page-1.png")
        self.assertEqual(path.name, "page-1.png")
        remove_images(user_id, md5_value)

        for filename in ("../secret.png", "page-0.png", "page-1.jpg"):
            with self.subTest(filename=filename):
                with self.assertRaises(HTTPException):
                    image_path(user_id, md5_value, filename)


if __name__ == "__main__":
    unittest.main()

"""从支持的文件中取得安全文件名和 (页号, 正文)。"""

from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document
from docx.table import Table
from pptx import Presentation

from app.core.failed_response import BusinessException
from app.services.document_parser import MAX_FILE_BYTES, parse_text_document

SUPPORTED_SUFFIXES = {".txt", ".md", ".pdf", ".docx", ".pptx"}


def extract_pages(filename: str, raw: bytes) -> tuple[str, list[tuple[int | None, str]]]:
    safe_name = filename.replace("\\", "/").split("/")[-1].strip()
    suffix = Path(safe_name).suffix.lower()
    if (
        not safe_name
        or len(safe_name) > 255
        or suffix not in SUPPORTED_SUFFIXES
    ):
        raise BusinessException(code=4001, message="文件类型不支持")
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise BusinessException(code=4001, message="文件必须非空且不超过 10 MB")
    if suffix in {".txt", ".md"}:
        _, text = parse_text_document(safe_name, raw)
        return safe_name, [(None, text)]

    try:
        if suffix == ".pdf":
            with pymupdf.open(stream=raw, filetype="pdf") as pdf:
                pages = [
                    (number + 1, page.get_text("text"))
                    for number, page in enumerate(pdf)
                ]
        elif suffix == ".docx":
            doc = Document(BytesIO(raw))
            parts = []
            for block in doc.iter_inner_content():
                if isinstance(block, Table):
                    parts.extend(
                        " | ".join(cell.text for cell in row.cells)
                        for row in block.rows
                    )
                else:
                    parts.append(block.text)
            pages = [(None, "\n".join(parts))]
        else:
            slides = Presentation(BytesIO(raw))
            pages = []
            for number, slide in enumerate(slides.slides, start=1):
                parts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text)
                    if shape.has_table:
                        parts.extend(
                            " | ".join(cell.text for cell in row.cells)
                            for row in shape.table.rows
                        )
                pages.append((number, "\n".join(parts)))
    except Exception as exc:
        raise BusinessException(code=4001, message="文件损坏或无法解析") from exc

    # 【修改】扫描 PDF 的空文字页可能有图，留给视觉服务处理。
    pages = [
        (number, text.strip())
        for number, text in pages
        if suffix == ".pdf" or text.strip()
    ]
    if not pages:
        raise BusinessException(code=4001, message="文件没有可提取文字")
    return safe_name, pages
